"""Resume text extraction and caching.

``extract_text`` turns an uploaded resume (PDF / DOCX / plain text) into plain
text. Optional parsers (``pypdf``, ``python-docx``) are imported lazily so a
deployment without them still works - it just yields less text.

``get_resume_text`` caches the result on ``jobs.CandidateProfile`` and only
re-parses when the underlying file changed.
"""

import hashlib
import io
import logging

from django.utils import timezone

logger = logging.getLogger(__name__)

MAX_CHARS = 12_000
MAX_BYTES = 5_000_000


def _read_bytes(file_field):
    try:
        file_field.open("rb")
        try:
            return file_field.read(MAX_BYTES)
        finally:
            try:
                file_field.close()
            except Exception:
                pass
    except Exception:
        logger.warning("Could not read resume file %r.", getattr(file_field, "name", "?"))
        return b""


def _pdf_text(raw):
    try:
        from pypdf import PdfReader
    except ImportError:
        logger.info("pypdf is not installed; skipping PDF resume text.")
        return ""
    try:
        reader = PdfReader(io.BytesIO(raw))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except Exception:
        logger.info("PDF resume could not be parsed.", exc_info=True)
        return ""


def _docx_text(raw):
    try:
        import docx
    except ImportError:
        logger.info("python-docx is not installed; skipping DOCX resume text.")
        return ""
    try:
        document = docx.Document(io.BytesIO(raw))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        return "\n".join(parts)
    except Exception:
        logger.info("DOCX resume could not be parsed.", exc_info=True)
        return ""


def _clean(text):
    lines = [line.strip() for line in (text or "").splitlines()]
    return "\n".join(line for line in lines if line)[:MAX_CHARS]


def extract_text(file_field, max_chars=MAX_CHARS):
    """Return plain text for an uploaded resume. Never raises."""
    if not file_field:
        return ""
    try:
        name = (getattr(file_field, "name", "") or "").lower()
        raw = _read_bytes(file_field)
        if not raw:
            return ""
        if name.endswith(".pdf") or raw[:5] == b"%PDF-":
            text = _pdf_text(raw)
        elif name.endswith((".docx", ".docm")) or raw[:2] == b"PK":
            text = _docx_text(raw)
        else:
            text = raw.decode("utf-8", errors="ignore")
        return _clean(text)[:max_chars]
    except Exception:  # pragma: no cover - extraction is strictly best-effort
        logger.warning("Resume text extraction failed.", exc_info=True)
        return ""


def resume_fingerprint(file_field):
    """Cheap identity for the stored file: name + size, hashed."""
    if not file_field:
        return ""
    name = getattr(file_field, "name", "") or ""
    try:
        size = file_field.size
    except Exception:
        size = ""
    return hashlib.sha256(f"{name}:{size}".encode()).hexdigest()


def get_resume_text(profile, force=False):
    """Return cached resume text for ``profile``, re-parsing when the file changed."""
    resume = getattr(profile, "resume", None)
    fingerprint = resume_fingerprint(resume)
    if not fingerprint:
        if profile.resume_text or profile.resume_hash:
            profile.resume_text = ""
            profile.resume_hash = ""
            profile.resume_parsed_at = None
            _save(profile)
        return ""
    if not force and fingerprint == profile.resume_hash and profile.resume_parsed_at:
        return profile.resume_text
    profile.resume_text = extract_text(resume)
    profile.resume_hash = fingerprint
    profile.resume_parsed_at = timezone.now()
    _save(profile)
    return profile.resume_text


def _save(profile):
    if profile.pk is None:
        return
    try:
        profile.save(update_fields=["resume_text", "resume_hash", "resume_parsed_at"])
    except Exception:  # pragma: no cover - caching must never break a request
        logger.warning("Could not cache resume text for profile %s.", profile.pk)
