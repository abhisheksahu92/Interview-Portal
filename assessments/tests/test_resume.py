"""Resume extraction/caching tests. Fixtures are generated, no binary assets."""

import io

import pytest
from django.core.files.base import ContentFile

from assessments import resume


def make_pdf_bytes(text="Senior Python engineer with Django experience"):
    """A tiny one-page PDF built with pypdf + a hand-written content stream."""
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)

    font = DictionaryObject()
    font[NameObject("/Type")] = NameObject("/Font")
    font[NameObject("/Subtype")] = NameObject("/Type1")
    font[NameObject("/BaseFont")] = NameObject("/Helvetica")
    font_ref = writer._add_object(font)
    fonts = DictionaryObject()
    fonts[NameObject("/F1")] = font_ref
    resources = DictionaryObject()
    resources[NameObject("/Font")] = fonts
    page[NameObject("/Resources")] = resources

    stream = DecodedStreamObject()
    stream.set_data(f"BT /F1 12 Tf 20 200 Td ({text}) Tj ET".encode())
    page[NameObject("/Contents")] = writer._add_object(stream)

    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def make_docx_bytes(text="Lead Django developer, 6 years"):
    import docx

    document = docx.Document()
    document.add_paragraph(text)
    document.add_paragraph("Skills: Python, PostgreSQL")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.fixture(autouse=True)
def _media(settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path


def test_extract_text_handles_missing_file():
    assert resume.extract_text(None) == ""
    assert resume.extract_text("") == ""


@pytest.mark.django_db
def test_extract_text_txt(candidate):
    candidate.resume.save("cv.txt", ContentFile(b"Backend engineer\n\n  Python  "), save=True)
    text = resume.extract_text(candidate.resume)
    assert "Backend engineer" in text
    assert "\n\n" not in text  # blank lines collapsed


@pytest.mark.django_db
def test_extract_text_pdf(candidate):
    candidate.resume.save("cv.pdf", ContentFile(make_pdf_bytes()), save=True)
    assert "Senior Python engineer" in resume.extract_text(candidate.resume)


@pytest.mark.django_db
def test_extract_text_docx(candidate):
    candidate.resume.save("cv.docx", ContentFile(make_docx_bytes()), save=True)
    text = resume.extract_text(candidate.resume)
    assert "Lead Django developer" in text
    assert "PostgreSQL" in text


@pytest.mark.django_db
def test_extract_text_caps_length(candidate):
    candidate.resume.save("cv.txt", ContentFile(b"x" * 50_000), save=True)
    assert len(resume.extract_text(candidate.resume)) == resume.MAX_CHARS


@pytest.mark.django_db
def test_extract_text_never_raises_on_broken_pdf(candidate):
    candidate.resume.save("cv.pdf", ContentFile(b"%PDF-1.4 not really a pdf"), save=True)
    assert resume.extract_text(candidate.resume) == ""


@pytest.mark.django_db
def test_missing_parser_degrades(candidate, monkeypatch):
    candidate.resume.save("cv.pdf", ContentFile(make_pdf_bytes()), save=True)
    import builtins

    real_import = builtins.__import__

    def no_pypdf(name, *args, **kwargs):
        if name == "pypdf":
            raise ImportError("no pypdf")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", no_pypdf)
    assert resume.extract_text(candidate.resume) == ""


@pytest.mark.django_db
def test_get_resume_text_caches_and_refreshes(candidate):
    assert resume.get_resume_text(candidate) == ""
    assert candidate.resume_parsed_at is None

    candidate.resume.save("cv.txt", ContentFile(b"First resume"), save=True)
    assert "First resume" in resume.get_resume_text(candidate)
    parsed_at = candidate.resume_parsed_at
    assert parsed_at is not None
    candidate.refresh_from_db()
    assert "First resume" in candidate.resume_text

    # Second call is served from the cache: no re-parse, timestamp unchanged.
    calls = []
    original = resume.extract_text
    resume.extract_text = lambda *a, **kw: calls.append(1) or ""
    try:
        assert "First resume" in resume.get_resume_text(candidate)
        assert calls == []
    finally:
        resume.extract_text = original

    # A new file changes the fingerprint and forces a re-parse.
    candidate.resume.save("cv2.txt", ContentFile(b"Second resume, longer"), save=True)
    assert "Second resume" in resume.get_resume_text(candidate)
    assert candidate.resume_parsed_at >= parsed_at


@pytest.mark.django_db
def test_get_resume_text_clears_cache_when_file_removed(candidate):
    candidate.resume.save("cv.txt", ContentFile(b"Some resume"), save=True)
    resume.get_resume_text(candidate)
    candidate.resume = ""
    assert resume.get_resume_text(candidate) == ""
    candidate.refresh_from_db()
    assert candidate.resume_text == ""
    assert candidate.resume_hash == ""
    assert candidate.resume_parsed_at is None


@pytest.mark.django_db
def test_force_reparses(candidate):
    candidate.resume.save("cv.txt", ContentFile(b"Cached"), save=True)
    resume.get_resume_text(candidate)
    candidate.resume_text = "stale"
    candidate.save(update_fields=["resume_text"])
    assert "Cached" in resume.get_resume_text(candidate, force=True)
